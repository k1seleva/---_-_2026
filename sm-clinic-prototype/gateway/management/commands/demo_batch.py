"""
Демо-поток для показа рабочего места: синтетические протоколы трёх клиник загружаются пачками (zip и файлы),
часть пациентов записывается в разные сроки, часть не приходит, по части врач завершает приём.

Только для разработки и автотестов интерфейса. В обычной установке не запускается: аналитика прототипа
показывает только реальные цифры (загруженные протоколы и действия пациентов, сводка прогона протоколов кейса).
Цифры, полученные после demo_batch, синтетические и для показа не годятся.
Запуск вручную: python manage.py demo_batch --i-know-it-is-synthetic
"""
import io
import random
import zipfile

from django.core.management.base import BaseCommand, CommandError
from docx import Document

from apps.doctors.models import Appointment
from apps.doctors.services.booking import BookingService, VisitCompletion, VisitService
from apps.patients.facade import PatientsFacade
from apps.patients.models import Patient
from apps.patients.services.booking import PatientBookingError, PatientBookingService
from apps.processing.models import UploadBatch
from apps.processing.services.pipeline import BatchIngestService, IncomingFile
from apps.routing.facade import RoutingFacade
from apps.routing.services.escalation import EscalationEngine
from common import clock

PELVIS = "Ультразвуковое исследование органов малого таза"
ABDOMEN = "УЗИ органов брюшной полости"
THYROID = "УЗИ щитовидной железы"
BREAST = "УЗИ молочных желез"
VEINS = "УЗДС вен нижних конечностей"
PROSTATE = "ТРУЗИ предстательной железы"
SOFT = "УЗИ мягких тканей передней брюшной стенки"

# (код сценария, тип исследования, строки описания, заключение или None, рекомендации или None)
SCENARIOS = [
    ("gallstones", ABDOMEN, ["ЖЕЛЧНЫЙ ПУЗЫРЬ: размеры 72*34 мм, стенка 2 мм.", "Конкременты множественные размером до 12 мм."],
     "УЗ-признаки холецистолитиаза.", "Консультация хирурга."),
    ("gb_polyp", ABDOMEN, ["ЖЕЛЧНЫЙ ПУЗЫРЬ: размеры обычные.", "На стенке пристеночное образование 11 мм, не смещается."],
     "Полип желчного пузыря 11 мм.", "Консультация хирурга."),
    ("endo_polyp", PELVIS, ["Матка: размеры 52х41х47 мм.", "М-эхо 9 мм, в полости гиперэхогенное включение 8 мм."],
     "Полип эндометрия.", "Консультация гинеколога."),
    ("myoma", PELVIS, ["Матка: размеры 61х50х55 мм.", "По задней стенке интрамуральный узел 24 мм."],
     "Миома матки (интрамуральный узел).", "Консультация гинеколога."),
    ("thyroid", THYROID, ["Правая доля: узловое образование 14 мм, гипоэхогенное, с неровным контуром."],
     "Узловое образование правой доли щитовидной железы, TI-RADS 4.", "Консультация эндокринолога, ТАБ."),
    ("breast", BREAST, ["Левая молочная железа: на 2 часах очаговое образование 12 мм с неровными контурами."],
     "Образование левой молочной железы, BI-RADS 4.", "Консультация маммолога."),
    ("hernia", SOFT, ["В области пупочного кольца дефект апоневроза 14 мм с выходом предбрюшинной клетчатки."],
     "Пупочная грыжа.", "Консультация хирурга."),
    ("bph", PROSTATE, ["Объём предстательной железы 48 см3, переходная зона увеличена."],
     "ДГПЖ.", "Консультация уролога."),
    ("varicose", VEINS, ["Несостоятельность ствола БПВ на бедре, рефлюкс 3 с."],
     "Варикозная трансформация БПВ слева.", "Консультация флеболога."),
    ("norm", ABDOMEN, ["ПЕЧЕНЬ: контуры ровные, структура однородная.", "ЖЕЛЧНЫЙ ПУЗЫРЬ: без особенностей."],
     "Патологических изменений не выявлено.", None),
    ("norm", PELVIS, ["Матка: размеры 48х38х42 мм, контуры ровные.", "Яичники обычных размеров."],
     "Эхографическая картина в пределах нормы.", None),
    # Неполные протоколы: попадут в «Нужна проверка» с причиной.
    ("no_recs", ABDOMEN, ["ЖЕЛЧНЫЙ ПУЗЫРЬ: стенка 3 мм.", "Конкременты до 9 мм."], "Холецистолитиаз.", None),
    ("no_conclusion", PELVIS, ["Матка: размеры 55х44х49 мм.", "Субмукозный миоматозный узел 15 мм, деформация полости матки."],
     None, None),
    ("uncertain", PELVIS, ["М-эхо 11 мм, неоднородное."], "Нельзя исключить полип эндометрия?", "Консультация гинеколога."),
    ("not_in_conclusion", ABDOMEN, ["ЖЕЛЧНЫЙ ПУЗЫРЬ: конкременты до 14 мм.", "ПЕЧЕНЬ: в правой доле анэхогенное образование 8 мм."],
     "УЗ-признаки холецистолитиаза.", "Консультация хирурга."),
]
EMERGENCY = ("dvt", VEINS, ["Просвет общей бедренной вены не сжимается, заполнен гипоэхогенными массами."],
             "Признаки окклюзивного тромбоза общей бедренной вены справа.", "Экстренная консультация сосудистого хирурга.")

FIRST_M = ["Алексей", "Дмитрий", "Игорь", "Павел", "Андрей", "Николай", "Виктор", "Роман", "Михаил", "Борис"]
FIRST_F = ["Светлана", "Ольга", "Людмила", "Галина", "Вера", "Юлия", "Ксения", "Валентина", "Дарья", "Нина"]
PATRONYMIC_M = ["Иванович", "Петрович", "Сергеевич", "Олегович", "Андреевич"]
PATRONYMIC_F = ["Ивановна", "Петровна", "Сергеевна", "Олеговна", "Андреевна"]
CLINICS = ["vdnh", "tekstilshchiki", "senezhskaya"]


def fits(scenario: tuple, sex: str) -> bool:
    """Сценарий подходит пациенту по полу: исследования малого таза и молочных желёз — женщинам, простаты — мужчинам."""
    code, study = scenario[0], scenario[1]
    if study == PELVIS or code == "breast":
        return sex != "M"
    if study == PROSTATE:
        return sex != "F"
    return True


def protocol_docx(card: str, date: str, study: str, lines: list[str], conclusion: str | None, recs: str | None) -> bytes:
    """Протокол в формате выгрузки 1С: шапка-таблица (карта, дата) и текст."""
    doc = Document()
    table = doc.add_table(rows=2, cols=4)
    head = table.rows[0].cells
    head[0].merge(head[1]).text = "Амбулаторная карта №"
    head[2].merge(head[3]).text = card
    cells = table.rows[1].cells
    cells[0].text, cells[1].text, cells[2].text, cells[3].text = "Дата приема:", date, "Время:", "10:00"
    doc.add_paragraph(study)
    doc.add_paragraph("Исследование выполнено на аппарате DEMO-SCAN.")
    for line in lines:
        doc.add_paragraph(line)
    if conclusion:
        doc.add_paragraph(f"ЗАКЛЮЧЕНИЕ: {conclusion}")
    doc.add_paragraph("Данное заключение не является клиническим диагнозом.")
    if recs:
        doc.add_paragraph(f"Рекомендовано: {recs}")
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


class Command(BaseCommand):
    """Синтетический поток. Требует явного флага, чтобы синтетика не попала в показ по ошибке."""

    help = "Синтетический поток протоколов трёх клиник и действия пациентов — для показа рабочего места"

    def add_arguments(self, parser):
        parser.add_argument("--patients", type=int, default=24, help="Сколько синтетических пациентов (кроме демо)")
        parser.add_argument("--seed", type=int, default=7)
        parser.add_argument("--i-know-it-is-synthetic", action="store_true", dest="confirm",
                            help="подтвердить, что нужны синтетические данные (только разработка и тесты)")

    def handle(self, *args, **opts):
        if not opts.get("confirm"):
            raise CommandError("demo_batch создаёт синтетические протоколы и записи. Для показа аналитики загрузите "
                               "настоящие протоколы. Если синтетика нужна для разработки: --i-know-it-is-synthetic")
        rnd = random.Random(opts["seed"])
        today = clock.now().date()
        patients = list(Patient.objects.filter(is_anonymous=False).exclude(external_mis_id__isnull=True))
        for i in range(opts["patients"]):
            female = i % 3 != 0
            # Свой счётчик для женщин и мужчин: имена не повторяются, отчество меняется на каждом круге имён.
            k = i - i // 3 - 1 if female else i // 3
            first, patronymic = (FIRST_F, PATRONYMIC_F) if female else (FIRST_M, PATRONYMIC_M)
            name = f"{first[k % len(first)]} {patronymic[(k + k // len(first)) % 5]}"
            pid = PatientsFacade.ensure_by_mis_id(f"AK-{1000 + i}", display_name=name,
                                                  birth_year=rnd.choice([1950, 1957, 1962, 1970, 1978, 1985, 1991]),
                                                  sex="F" if female else "M")
            patients.append(Patient.objects.get(pk=pid))

        per_clinic: dict[str, list[IncomingFile]] = {c: [] for c in CLINICS}
        pointer = 0
        for n, patient in enumerate(patients):
            # Сценарии по кругу, но только подходящие по полу (миома — не мужчине, простата — не женщине).
            index = next(i for i in range(pointer, pointer + len(SCENARIOS)) if fits(SCENARIOS[i % len(SCENARIOS)], patient.sex))
            pointer = index + 1
            code, study, lines, conclusion, recs = SCENARIOS[index % len(SCENARIOS)]
            clinic = CLINICS[n % 3]
            date = (today.replace(day=1) if today.day > 3 else today).strftime("%d.%m.%Y")
            data = protocol_docx(patient.external_mis_id, date, study, lines, conclusion, recs)
            per_clinic[clinic].append(IncomingFile(name=f"{patient.external_mis_id}.docx", data=data))
        # Протоколы без узнаваемого номера карты: попадут в «Обезличенные», один из них экстренный.
        unknown = [("К-7781", SCENARIOS[0]), ("К-7781", SCENARIOS[4]), ("", SCENARIOS[2]), ("Б/Н 15", EMERGENCY)]
        for i, (card, (code, study, lines, conclusion, recs)) in enumerate(unknown):
            per_clinic[CLINICS[i % 3]].append(IncomingFile(
                name=f"протокол-{i + 1}.docx",
                data=protocol_docx(card, today.strftime("%d.%m.%Y"), study, lines, conclusion, recs)))
        # Экстренный протокол известного пациента и испорченный файл.
        code, study, lines, conclusion, recs = EMERGENCY
        per_clinic["vdnh"].append(IncomingFile(name="AK-0007-УЗДС.docx", data=protocol_docx(
            "AK-0007", today.strftime("%d.%m.%Y"), study, lines, conclusion, recs)))
        per_clinic["tekstilshchiki"].append(IncomingFile(name="скан-повреждён.docx", data=b"PK\x03\x04 not a docx"))

        service = BatchIngestService()
        for clinic, files in per_clinic.items():
            if clinic == "senezhskaya":
                report = service.ingest(files, location_code=clinic, source=UploadBatch.Source.MANUAL, created_by="demo")
            else:
                buf = io.BytesIO()
                with zipfile.ZipFile(buf, "w") as zf:
                    for f in files:
                        zf.writestr(f"Протоколы/{f.name}", f.data)
                    zf.writestr("__MACOSX/._junk", b"")
                report = service.ingest([IncomingFile(name=f"{clinic}.zip", data=buf.getvalue())], location_code=clinic,
                                        created_by="demo")
            self.stdout.write(f"{clinic}: принято {report.batch.accepted}, не принято {len(report.batch.rejected)}")
        # Повторная загрузка того же файла — не дубль.
        service.ingest(per_clinic["senezhskaya"][:1], location_code="senezhskaya", created_by="demo")

        booked = self._book(rnd, share=0.45)                         # в первые сутки
        self._advance(48)
        booked += self._book(rnd, share=0.35)                        # к 72 часам
        self._advance(24 * 6)                                        # напоминания, звонки координатора
        booked += self._book(rnd, share=0.3)
        self._visits(rnd)
        self.stdout.write(self.style.SUCCESS(f"Демо-поток готов: записей {booked}, модельное время {clock.now():%d.%m.%Y %H:%M}"))

    @staticmethod
    def _advance(hours: int) -> None:
        clock.advance(hours=hours)
        EscalationEngine().tick()

    @staticmethod
    def _book(rnd: random.Random, share: float) -> int:
        service, count = PatientBookingService(), 0
        for patient in Patient.objects.filter(is_anonymous=False):
            for route in RoutingFacade.list_patient_routes(patient.id, open_only=True):
                step = route.get("active_step")
                if not step or step["status"] != "awaiting_booking" or rnd.random() > share:
                    continue
                try:
                    _, slots = service.available_slots(patient, step["id"])
                    if slots:
                        service.book(patient, step["id"], slots[min(len(slots) - 1, rnd.randint(0, 3))]["id"])
                        count += 1
                except PatientBookingError:
                    continue
        return count

    @staticmethod
    def _visits(rnd: random.Random) -> None:
        """Часть пациентов не пришла, по части врач завершил приём с тактикой."""
        tactics = ["surgery_indicated", "surgery_indicated", "observation", "extra_exam", "surgery_not_indicated"]
        for appt in Appointment.objects.filter(status__in=[Appointment.Status.SCHEDULED, Appointment.Status.CONFIRMED]).select_related("slot__doctor"):
            roll = rnd.random()
            if roll < 0.15:
                BookingService().mark_no_show(appt)
            elif roll < 0.6:
                tactic = rnd.choice(tactics)
                agrees = tactic != "surgery_not_indicated"
                VisitService().complete(appt, appt.slot.doctor, VisitCompletion(
                    tactic=tactic, agrees_with_ai_route=agrees,
                    disagreement_reason="" if agrees else "По клинической картине операция не нужна"))
