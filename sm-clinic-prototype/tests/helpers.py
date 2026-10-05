"""Общие помощники тестов: синтетические протоколы (выданные файлы в репозиторий не кладём)."""
from io import BytesIO

from django.core.management import call_command
from docx import Document


def make_docx(conclusion: str, title: str = "УЛЬТРАЗВУКОВОЕ ИССЛЕДОВАНИЕ ОРГАНОВ МАЛОГО ТАЗА", date: str = "26.08.2026") -> bytes:
    """Протокол в формате выгрузки 1С: шапка-таблица с объединёнными ячейками + текст."""
    doc = Document()
    table = doc.add_table(rows=2, cols=4)
    row = table.rows[0].cells
    row[0].merge(row[1]).text = "Амбулаторная карта №"
    row[2].merge(row[3]).text = "TEST-1"
    cells = table.rows[1].cells
    cells[0].text, cells[1].text, cells[2].text, cells[3].text = "Дата приема:", date, "Время:", "10:00"
    doc.add_paragraph(title)
    doc.add_paragraph("Матка: размеры 50х40х45 мм. Объемные образования яичников: не лоцируются.")
    doc.add_paragraph(f"ЗАКЛЮЧЕНИЕ: {conclusion}")
    doc.add_paragraph("Данное заключение не является клиническим диагнозом.")
    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()


def make_protocol_docx(lines: list[str], date: str = "26.08.2026") -> bytes:
    """Протокол 1С с произвольным описанием (строки — абзацы после шапки)."""
    doc = Document()
    table = doc.add_table(rows=2, cols=4)
    row = table.rows[0].cells
    row[0].merge(row[1]).text = "Амбулаторная карта №"
    row[2].merge(row[3]).text = "TEST-2"
    cells = table.rows[1].cells
    cells[0].text, cells[1].text, cells[2].text, cells[3].text = "Дата приема:", date, "Время:", "10:00"
    for line in lines:
        doc.add_paragraph(line)
    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()


# Синтетический протокол УЗИ желчного пузыря: конкременты в описании, ЖКБ в заключении,
# а рекомендация — «лечащий врач» без профиля; в печени образование, не вынесенное в заключение.
GALLSTONE_PROTOCOL = [
    "УЗИ органов брюшной полости",
    "Исследование выполнено на аппарате TEST-SCAN конвексным датчиком.",
    "ЖЕЛЧНЫЙ ПУЗЫРЬ: Расположение обычное. Размеры не увеличены 70*35 мм. Стенка 2 мм, не утолщена.",
    "Конкременты множественные размером до 14 мм.",
    "ПЕЧЕНЬ",
    "Контуры ровные, четкие. Очаговых образований не выявлено.",
    "В правой доле лоцируется анэхогенное образование 8 мм.",
    "ЗАКЛЮЧЕНИЕ: УЗ-признаки холецистолитиаза.",
    "Данное заключение не является диагнозом.",
    "Рекомендовано: консультация лечащего врача.",
]


def seed() -> None:
    # Предохранитель модели хранится в памяти процесса: сбой модели в одном тесте не должен влиять на другой.
    from apps.processing.services.ai_agent import reset_circuits

    reset_circuits()
    call_command("seed_demo", days=40, verbosity=0)


def staff_client(role: str = "coordinator", username: str = ""):
    """Клиент, вошедший как сотрудник через страницу входа (аккаунты создаёт seed_demo)."""
    from django.conf import settings
    from django.test import Client

    from django.core.cache import cache

    client = Client()
    login = username or {"coordinator": "coordinator", "doctor": "doctor1"}[role]
    cache.delete(f"staff-login:{login}")  # блокировку после неверных паролей из других тестов не наследуем
    response = client.post(f"/{role}/login/", {"username": login, "password": settings.STAFF_LOGIN["DEMO_PASSWORD"]})
    assert response.status_code == 302, response.content[:300]
    return client
